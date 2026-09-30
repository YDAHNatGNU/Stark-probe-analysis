# Supplementary Software for "Solvent-Reconstructed Electric Fields at
# Solid-Liquid Interfaces" (Cho et al.). MIT License; see LICENSE.
"""Differential cumulant analysis of the Stark line shape. The observed line is
the convolution of the intrinsic emission with the distribution of Stark shifts,
and cumulants are additive under convolution, so the cumulants of the Stark
kernel follow by subtracting those of the Stark-free control. No line-shape model
is assumed."""
import numpy as np
import openpyxl
import os as _os
# Input data are read from the data folder beside the analysis folders;
# the location can be overridden with the DATA_DIR environment variable.
_DATA = _os.environ.get("DATA_DIR") or _os.path.join(
    _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))), "data")

XLSX = _os.path.join(_DATA, "Spectrum_Rawdata_tBTO_SiO2_in_EtOH.xlsx")
W     = 0.080      # window half-width (eV)
NBOOT = 20_000
rng   = np.random.default_rng(0)


# ---------------------------------------------------------------- data
def load(path):
    """Return {sheet_name: (E, Y)} with E in eV and Y of shape (n_points, n_spectra)."""
    wb, out = openpyxl.load_workbook(path, data_only=True), {}
    for ws in wb:
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        E = np.array([r[0] for r in rows], float)
        Y = np.array([[r[j] for j in range(1, ws.max_column) if r[j] is not None]
                      for r in rows], float)
        o = np.argsort(E)
        out[ws.title] = (E[o], Y[o])
    return out


# ------------------------------------------------------------ estimator
def cumulants(E, y, lo, hi):
    """Eq. (SX.6) on one spectrum over the window [lo, hi]."""
    sel  = (E >= lo) & (E <= hi)
    x, v = E[sel], y[sel].astype(float)
    bg   = np.median(np.r_[y[E < lo][-25:], y[E > hi][:25]])   # local background
    v    = np.clip(v - bg, 0, None)                            # no negative weights
    w    = v / v.sum()                                         # unit area
    k1 = np.sum(w * x)
    k2 = np.sum(w * (x - k1) ** 2)
    k3 = np.sum(w * (x - k1) ** 3)
    return k1, k2, k3


def per_spectrum(E, Y, w=W, min_signal=5000):
    """Cumulants of every spectrum, window centred on its own peak."""
    out = []
    for j in range(Y.shape[1]):
        y  = Y[:, j].astype(float)
        yb = y - np.median(y[:30])
        if yb.max() < min_signal:                              # skip dim spectra
            continue
        pk = E[int(np.argmax(yb))]
        out.append(cumulants(E, y, pk - w, pk + w))
    return np.array(out)                                       # (n_spectra, 3)


# ------------------------------------------------------------ bootstrap
def bootstrap(a, b, n=NBOOT):
    """Sampling distributions of dkappa2, dkappa3, gamma1 and |m|."""
    K2 = np.empty(n); K3 = np.empty(n)
    for i in range(n):
        ai = rng.integers(0, len(a), len(a))                   # with replacement
        bi = rng.integers(0, len(b), len(b))
        K2[i] = a[ai, 1].mean() - b[bi, 1].mean()
        K3[i] = a[ai, 2].mean() - b[bi, 2].mean()
    ok = K2 > 0
    G  = K3[ok] / K2[ok] ** 1.5                                # Eq. (SX.6): gamma1
    M  = np.abs(G) / np.sqrt(4 + G ** 2)                       # Eq. (SX.7): |m|
    return K2, K3, G, M


# ----------------------------------------------------------------- main
def report(K2, K3, G, M, label=""):
    f = lambda v: (v.mean(), v.std(), v.mean() / v.std())
    print(f"\n--- {label} ---")
    print("dkappa2 = {:7.1f} +- {:5.1f} meV^2   (t = {:5.2f})".format(*[x * s for x, s
          in zip(f(K2), (1e6, 1e6, 1))]))
    print("dkappa3 = {:7.1f} +- {:5.1f} meV^3   (t = {:5.2f})".format(*[x * s for x, s
          in zip(f(K3), (1e9, 1e9, 1))]))
    lo, hi = np.percentile(G, [2.5, 97.5])
    print(f"gamma1  = {np.median(G):.3f}   95% CI [{lo:.3f}, {hi:.3f}]")
    print(f"|m|     = {np.median(M):.3f}   95% upper bound {np.percentile(M, 95):.3f}")


if __name__ == "__main__":
    D = load(XLSX)
    a = per_spectrum(*D["tBTO"])       # shell-free tBTO, d = 0
    b = per_spectrum(*D["SiO2"])       # Stark-free control
    print(f"tBTO {len(a)} spectra, SiO2 {len(b)} spectra")
    report(*bootstrap(a, b), label=f"W = {W*1000:.0f} meV (primary)")

    # window scan -> Table S1.3
    # Each window uses the primary estimator: 20,000 resamples, generator
    # re-seeded (seed 0) per window, so the W = 80 meV row equals the primary result.
    for w in (0.06, 0.07, 0.08, 0.09, 0.10, 0.12, 0.15):
        rng = np.random.default_rng(0)
        aa, bb = per_spectrum(*D["tBTO"], w=w), per_spectrum(*D["SiO2"], w=w)
        report(*bootstrap(aa, bb, n=NBOOT), label=f"W = {w*1000:.0f} meV")
